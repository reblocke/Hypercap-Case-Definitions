#!/usr/bin/env python3
"""Compare one legacy baseline and two guarded Stata runs without exposing values."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import numbers
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path, PureWindowsPath
from typing import Any, Iterable

from openpyxl import load_workbook
from PIL import Image

import input_manifest
import stata_run

ABS_TOL = 1e-12
REL_TOL = 1e-12
ROOT = Path(__file__).resolve().parents[1]
KNOWN_OUTPUT_RELATIVE_PATHS = tuple(
    sorted(
        {
            *stata_run.EXPECTED_XLSX,
            *stata_run.EXPECTED_PNG,
            *stata_run.LEGACY_PNG_ALIASES.values(),
            *(f"graph-temp/{name}" for name in stata_run.EXPECTED_GPH),
        },
        key=len,
        reverse=True,
    )
)
CORRECTION_MODE_INTEGRITY_CATEGORIES = frozenset(
    {
        "analysis_log_output_roots",
    }
)


class ArtifactRootFailure(ValueError):
    """Raised when a manifest artifact root is unsafe or unavailable."""


class LogOutputRootFailure(ValueError):
    """Raised when a transcript reports expected artifacts under multiple roots."""


def digest_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def resolve_artifact_root(path: Path, manifest: dict[str, Any]) -> Path:
    relative = manifest.get("artifact_relative_path", ".")
    if not isinstance(relative, str):
        raise ArtifactRootFailure
    relative = relative or "."
    relative_path = Path(relative)
    windows_path = PureWindowsPath(relative)
    if (
        relative_path.is_absolute()
        or windows_path.is_absolute()
        or windows_path.drive
        or ".." in relative_path.parts
        or ".." in windows_path.parts
    ):
        raise ArtifactRootFailure
    try:
        run_root = path.resolve(strict=True)
        artifact_root = (run_root / relative_path).resolve(strict=True)
    except OSError as exc:
        raise ArtifactRootFailure from exc
    if not artifact_root.is_dir() or not artifact_root.is_relative_to(run_root):
        raise ArtifactRootFailure
    return artifact_root


def load_run(path: Path) -> tuple[dict[str, Any], Path]:
    manifest_path = path / "run_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    artifact_root = resolve_artifact_root(path, manifest)
    return manifest, artifact_root


def compare_scalar(
    left: Any,
    right: Any,
    *,
    artifact: str,
    location: str,
    failures: list[dict[str, Any]],
) -> None:
    if (
        isinstance(left, numbers.Real)
        and not isinstance(left, bool)
        and isinstance(right, numbers.Real)
        and not isinstance(right, bool)
    ):
        if math.isclose(float(left), float(right), rel_tol=REL_TOL, abs_tol=ABS_TOL):
            return
    elif left == right:
        return
    failures.append(
        {
            "category": "workbook_value",
            "artifact": artifact,
            "location": location,
        }
    )


def compare_workbook(
    left_path: Path,
    right_path: Path,
    failures: list[dict[str, Any]],
) -> None:
    artifact = left_path.name
    left_book = load_workbook(left_path, data_only=False, read_only=False)
    right_book = load_workbook(right_path, data_only=False, read_only=False)
    try:
        if left_book.sheetnames != right_book.sheetnames:
            failures.append({"category": "workbook_sheets", "artifact": artifact})
            return
        for sheet_name in left_book.sheetnames:
            left_sheet = left_book[sheet_name]
            right_sheet = right_book[sheet_name]
            if (left_sheet.max_row, left_sheet.max_column) != (
                right_sheet.max_row,
                right_sheet.max_column,
            ):
                failures.append(
                    {
                        "category": "workbook_dimensions",
                        "artifact": artifact,
                        "location": sheet_name,
                    }
                )
                continue
            left_merges = sorted(str(item) for item in left_sheet.merged_cells.ranges)
            right_merges = sorted(str(item) for item in right_sheet.merged_cells.ranges)
            if left_merges != right_merges:
                failures.append(
                    {
                        "category": "workbook_merges",
                        "artifact": artifact,
                        "location": sheet_name,
                    }
                )
            for row in left_sheet.iter_rows():
                for left_cell in row:
                    right_cell = right_sheet[left_cell.coordinate]
                    compare_scalar(
                        left_cell.value,
                        right_cell.value,
                        artifact=artifact,
                        location=f"{sheet_name}!{left_cell.coordinate}",
                        failures=failures,
                    )
                    if (
                        left_cell.data_type != right_cell.data_type
                        or left_cell.number_format != right_cell.number_format
                    ):
                        failures.append(
                            {
                                "category": "workbook_cell_metadata",
                                "artifact": artifact,
                                "location": f"{sheet_name}!{left_cell.coordinate}",
                            }
                        )
    finally:
        left_book.close()
        right_book.close()


def compare_png(
    left_path: Path,
    right_path: Path,
    failures: list[dict[str, Any]],
    *,
    artifact: str | None = None,
) -> None:
    artifact = artifact or left_path.name
    with Image.open(left_path) as left_image, Image.open(right_path) as right_image:
        left = left_image.convert("RGBA")
        right = right_image.convert("RGBA")
        if left.size != right.size:
            failures.append({"category": "png_dimensions", "artifact": artifact})
            return
        if left.tobytes() != right.tobytes():
            failures.append({"category": "png_pixels", "artifact": artifact})


def strip_row_listing(lines: list[str]) -> list[str]:
    output: list[str] = []
    skipping = False
    for line in lines:
        if "list hypercap_resp_failure paco2 vbg_co2" in line:
            skipping = True
            continue
        if skipping and "/* Ones I can't do */" in line:
            skipping = False
            output.append(line)
            continue
        if not skipping:
            output.append(line)
    return output


def unwrap_stata_continuations(lines: list[str]) -> list[str]:
    output: list[str] = []
    for line in lines:
        stripped = line.lstrip()
        if stripped.startswith("> ") and output:
            output[-1] += stripped[2:]
        else:
            output.append(line)
    return output


def unwrap_stata_parenthesized_file_messages(lines: list[str]) -> list[str]:
    output: list[str] = []
    index = 0
    while index < len(lines):
        message = lines[index]
        if not message.startswith("(file ") or message.rstrip().endswith(")"):
            output.append(message)
            index += 1
            continue
        while index + 1 < len(lines) and not message.rstrip().endswith(")"):
            message += " " + lines[index + 1].strip()
            index += 1
        output.append(message)
        index += 1
    return output


def unwrap_stata_file_messages(lines: list[str]) -> list[str]:
    output: list[str] = []
    index = 0
    while index < len(lines):
        message = lines[index]
        if not message.startswith("file "):
            output.append(message)
            index += 1
            continue

        while index + 1 < len(lines):
            following = lines[index + 1]
            missing_saved = " saved" not in message
            incomplete_format = (
                " saved as" in message
                and not message.rstrip().endswith(" format")
            )
            wrapped_saved_as = (
                message.rstrip().endswith(" saved")
                and following.lstrip().startswith("as ")
            )
            if not (missing_saved or incomplete_format or wrapped_saved_as):
                break
            message += " " + following.strip()
            index += 1
        output.append(message)
        index += 1
    return output


def redact_wrapped_path(text: str, path: Path, replacement: str) -> str:
    value = path.as_posix()
    gap = r"(?:\r?\n[ \t]*)?"
    characters = [
        r"(?:[ \t]+|\r?\n[ \t]*)" if character == " " else re.escape(character)
        for character in value
    ]
    pattern = gap.join(characters)
    return re.sub(pattern, replacement, text)


def redact_known_output_paths(text: str) -> str:
    """Redact relocated roots in Stata's expected-artifact file notifications."""
    patterns = (
        re.compile(
            r"^(?P<prefix>file )(?P<path>.+)"
            r"(?P<suffix> saved(?: as [^\r\n]+ format)?)$"
        ),
        re.compile(
            r"^(?P<prefix>\(file )(?P<path>.+)(?P<suffix> not found\))$"
        ),
    )
    lines = text.split("\n")
    notifications: dict[int, tuple[re.Match[str], str]] = {}
    observed_roots: set[str] = set()
    for index, line in enumerate(lines):
        match: re.Match[str] | None = None
        for pattern in patterns:
            match = pattern.fullmatch(line)
            if match is not None:
                break
        if match is None:
            continue
        observed_path = match.group("path")
        normalized_path = observed_path.replace("\\", "/")
        if not (
            normalized_path.startswith("/")
            or re.match(r"^[A-Za-z]:/", normalized_path)
        ):
            continue
        relative = next(
            (
                candidate
                for candidate in KNOWN_OUTPUT_RELATIVE_PATHS
                if normalized_path.endswith(f"/{candidate}")
            ),
            None,
        )
        if relative is None:
            continue
        observed_roots.add(normalized_path[: -(len(relative) + 1)])
        notifications[index] = (match, relative)
    if len(observed_roots) > 1:
        raise LogOutputRootFailure
    if not observed_roots:
        return text
    for index, (match, relative) in notifications.items():
        lines[index] = (
            f"{match.group('prefix')}<OUTPUT>/{relative}{match.group('suffix')}"
        )
    return "\n".join(lines)


def normalize_log(path: Path, artifact_root: Path, run_root: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    start_match = re.search(r"(?m)^.*Pre-processing\s*$", text)
    if not start_match:
        raise ValueError("preprocessing marker unavailable")
    start = text.rfind("/*", 0, start_match.start())
    if start < 0:
        start = start_match.start()

    end_candidates = [
        location
        for location in (
            text.find("/* Aggregate-only run metrics", start),
            text.find("\nend of do-file", start),
        )
        if location >= 0
    ]
    end = min(end_candidates) if end_candidates else len(text)
    lines = unwrap_stata_continuations(text[start:end].splitlines())
    lines = unwrap_stata_parenthesized_file_messages(lines)
    lines = unwrap_stata_file_messages(lines)
    excerpt = "\n".join(lines)
    excerpt = redact_known_output_paths(excerpt)
    roots = sorted(
        {artifact_root, run_root, artifact_root.parent},
        key=lambda item: len(item.as_posix()),
        reverse=True,
    )
    for root in roots:
        excerpt = redact_wrapped_path(excerpt, root, "<OUTPUT>")
    mac_user_prefix = "/" + "Users" + "/"
    excerpt = re.sub(
        re.escape(mac_user_prefix) + r"[^\n\"]+/full_db[.]dta",
        "<INPUT>/full_db.dta",
        excerpt,
    )
    lines = strip_row_listing(excerpt.splitlines())
    normalized: list[str] = []
    previous_blank = False
    for line in lines:
        clean = line.rstrip()
        if clean.strip() == ".":
            continue
        blank = not clean
        if blank and previous_blank:
            continue
        normalized.append(clean)
        previous_blank = blank
    return "\n".join(normalized).strip() + "\n"


def compare_logs(
    left_manifest: dict[str, Any],
    left_artifacts: Path,
    left_run: Path,
    right_manifest: dict[str, Any],
    right_artifacts: Path,
    right_run: Path,
    failures: list[dict[str, Any]],
) -> None:
    left_log = left_artifacts / left_manifest["artifacts"]["analysis_log"]
    right_log = right_artifacts / right_manifest["artifacts"]["analysis_log"]
    normalized_logs: dict[str, str] = {}
    for role, log_path, artifact_root, run_root in (
        ("left", left_log, left_artifacts, left_run),
        ("right", right_log, right_artifacts, right_run),
    ):
        try:
            normalized_logs[role] = normalize_log(
                log_path,
                artifact_root,
                run_root,
            )
        except LogOutputRootFailure:
            failures.append(
                {
                    "category": "analysis_log_output_roots",
                    "artifact": "analysis_log",
                    "role": role,
                }
            )
    if len(normalized_logs) != 2:
        return
    left_text = normalized_logs["left"]
    right_text = normalized_logs["right"]
    if left_text == right_text:
        return
    left_lines = left_text.splitlines()
    right_lines = right_text.splitlines()
    differing_line = min(len(left_lines), len(right_lines)) + 1
    for index, (left, right) in enumerate(zip(left_lines, right_lines), start=1):
        if left != right:
            differing_line = index
            break
    failures.append(
        {
            "category": "analysis_log",
            "artifact": "analysis_log",
            "normalized_line": differing_line,
            "left_sha256": digest_text(left_text),
            "right_sha256": digest_text(right_text),
        }
    )


def compare_environment(
    left: dict[str, Any],
    right: dict[str, Any],
    failures: list[dict[str, Any]],
) -> None:
    if left["input"]["sha256"] != right["input"]["sha256"]:
        failures.append({"category": "input_hash", "artifact": "run_manifest"})
    left_status = left.get("driver_status", {})
    right_status = right.get("driver_status", {})
    for key in ("stata_version", "stata_flavor", "operating_system", "machine_type"):
        if left_status.get(key) != right_status.get(key):
            failures.append(
                {
                    "category": "stata_environment",
                    "artifact": "run_manifest",
                    "field": key,
                }
            )
    left_dependencies = {
        row["name"]: (row["found"], row["sha256"], row["version_header"])
        for row in left.get("dependencies", [])
    }
    right_dependencies = {
        row["name"]: (row["found"], row["sha256"], row["version_header"])
        for row in right.get("dependencies", [])
    }
    if left_dependencies != right_dependencies:
        failures.append({"category": "dependencies", "artifact": "dependency_report"})


def compare_pair(
    left_run: Path,
    right_run: Path,
    label: str,
) -> dict[str, Any]:
    left_manifest, left_artifacts = load_run(left_run)
    right_manifest, right_artifacts = load_run(right_run)
    failures: list[dict[str, Any]] = []

    compare_environment(left_manifest, right_manifest, failures)
    for manifest in (left_manifest, right_manifest):
        if manifest.get("status") != "success":
            failures.append({"category": "run_status", "artifact": "run_manifest"})
        if manifest.get("analysis_worktree_dirty"):
            failures.append({"category": "dirty_worktree", "artifact": "run_manifest"})
        if manifest.get("artifacts", {}).get("missing"):
            failures.append({"category": "artifact_inventory", "artifact": "run_manifest"})
        if manifest.get("artifacts", {}).get("control_missing"):
            failures.append({"category": "control_inventory", "artifact": "run_manifest"})

    for name in stata_run.EXPECTED_XLSX:
        compare_workbook(left_artifacts / name, right_artifacts / name, failures)
    for name in stata_run.EXPECTED_PNG:
        left_name = (
            stata_run.LEGACY_PNG_ALIASES.get(name, name)
            if left_manifest.get("legacy_two_argument_mode") is True
            else name
        )
        right_name = (
            stata_run.LEGACY_PNG_ALIASES.get(name, name)
            if right_manifest.get("legacy_two_argument_mode") is True
            else name
        )
        compare_png(
            left_artifacts / left_name,
            right_artifacts / right_name,
            failures,
            artifact=name,
        )
    for name in stata_run.EXPECTED_GPH:
        for root in (left_artifacts, right_artifacts):
            path = root / "graph-temp" / name
            if not path.is_file() or path.stat().st_size == 0:
                failures.append({"category": "gph_presence", "artifact": name})

    for manifest, artifact_root in (
        (left_manifest, left_artifacts),
        (right_manifest, right_artifacts),
    ):
        copied = artifact_root / manifest["artifacts"]["copied_do"]
        expected = manifest["code_hashes"][f"analysis:{stata_run.MAIN_DO}"]
        if stata_run.sha256_file(copied) != expected:
            failures.append({"category": "copied_do_hash", "artifact": copied.name})

    compare_logs(
        left_manifest,
        left_artifacts,
        left_run,
        right_manifest,
        right_artifacts,
        right_run,
        failures,
    )
    return {
        "comparison": label,
        "status": "pass" if not failures else "fail",
        "failures": failures,
    }


def validate_run_set(
    runs: dict[str, Path],
    manifests: dict[str, dict[str, Any]],
    artifact_roots: dict[str, Path],
) -> list[dict[str, str]]:
    failures: list[dict[str, str]] = []
    run_paths = list(runs.values())
    duplicate_run = len(set(run_paths)) != len(run_paths)
    if not duplicate_run:
        for index, left in enumerate(run_paths):
            for right in run_paths[index + 1 :]:
                try:
                    if left.samefile(right):
                        duplicate_run = True
                        break
                except OSError:
                    continue
            if duplicate_run:
                break
    if duplicate_run:
        failures.append({"category": "run_isolation", "role": "run_set"})

    artifact_paths = list(artifact_roots.values())
    duplicate_artifacts = len(set(artifact_paths)) != len(artifact_paths)
    if not duplicate_artifacts:
        for index, left in enumerate(artifact_paths):
            for right in artifact_paths[index + 1 :]:
                try:
                    if left.samefile(right):
                        duplicate_artifacts = True
                        break
                except OSError:
                    continue
            if duplicate_artifacts:
                break
    if duplicate_artifacts:
        failures.append({"category": "artifact_isolation", "role": "run_set"})

    expected_legacy_modes = {
        "baseline": True,
        "candidate_1": False,
        "candidate_2": False,
    }
    for role, expected in expected_legacy_modes.items():
        if manifests[role].get("legacy_two_argument_mode") is not expected:
            failures.append({"category": "run_role", "role": role})

    candidate_commits: dict[str, str] = {}
    for role in ("candidate_1", "candidate_2"):
        commit = manifests[role].get("analysis_commit")
        if not isinstance(commit, str) or not commit.strip():
            failures.append({"category": "analysis_commit", "role": role})
            continue
        candidate_commits[role] = commit.strip().lower()
    if (
        len(candidate_commits) == 2
        and candidate_commits["candidate_1"] != candidate_commits["candidate_2"]
    ):
        failures.append(
            {"category": "candidate_commit_match", "role": "candidate_set"}
        )

    candidate_run_ids: dict[str, str] = {}
    for role in ("candidate_1", "candidate_2"):
        run_id = manifests[role].get("run_id")
        if not isinstance(run_id, str) or not run_id.strip():
            failures.append({"category": "run_id", "role": role})
            continue
        candidate_run_ids[role] = run_id.strip()
    if (
        len(candidate_run_ids) == 2
        and candidate_run_ids["candidate_1"] == candidate_run_ids["candidate_2"]
    ):
        failures.append(
            {"category": "candidate_run_id_match", "role": "candidate_set"}
        )
    return failures


def validate_live_run_evidence(
    runs: dict[str, Path],
    manifests: dict[str, dict[str, Any]],
    artifact_roots: dict[str, Path],
) -> list[dict[str, str]]:
    failures: list[dict[str, str]] = []
    for role, run_dir in runs.items():
        manifest = manifests[role]
        artifact_root = artifact_roots[role]
        legacy = manifest.get("legacy_two_argument_mode") is True
        inventory = stata_run.validate_artifact_inventory(
            artifact_root,
            legacy=legacy,
        )
        if inventory["missing"]:
            failures.append({"category": "artifact_inventory", "role": role})
        controls = stata_run.validate_control_artifacts(
            run_dir,
            artifact_root,
            inventory,
            legacy=legacy,
            include_success=True,
        )
        if controls["control_missing"]:
            failures.append({"category": "control_inventory", "role": role})
    return failures


def validate_post_run_input(
    observed_input_hash: str,
    manifests: dict[str, dict[str, Any]],
    expected_input_hash: str | None,
) -> list[dict[str, str]]:
    failures: list[dict[str, str]] = []
    for role, manifest in manifests.items():
        recorded_hash = manifest.get("input", {}).get("sha256")
        if (
            not isinstance(recorded_hash, str)
            or observed_input_hash.lower() != recorded_hash.lower()
        ):
            failures.append({"category": "post_run_input_hash", "role": role})
    if (
        expected_input_hash
        and observed_input_hash.lower() != expected_input_hash.lower()
    ):
        failures.append({"category": "post_run_input_hash", "role": "expected"})
    return failures


def validate_run_input_manifest_references(
    manifests: dict[str, dict[str, Any]],
    current_approval: dict[str, Any],
) -> list[dict[str, str]]:
    failures: list[dict[str, str]] = []
    for role in ("candidate_1", "candidate_2"):
        input_record = manifests[role].get("input", {})
        if (
            not isinstance(input_record, dict)
            or "approval" not in input_record
            or input_record["approval"] != current_approval
        ):
            failures.append(
                {"category": "run_input_manifest_reference", "role": role}
            )

    baseline_input = manifests["baseline"].get("input", {})
    if (
        isinstance(baseline_input, dict)
        and "approval" in baseline_input
        and baseline_input["approval"] != current_approval
    ):
        failures.append(
            {"category": "run_input_manifest_reference", "role": "baseline"}
        )
    return failures


def write_report(
    report_path: Path,
    *,
    observed_input_hash: str | None,
    comparisons: list[dict[str, Any]],
    failures: list[dict[str, Any]],
    comparison_mode: str = "equivalence",
    announce: bool = True,
) -> dict[str, Any]:
    by_label = {
        item.get("comparison"): item
        for item in comparisons
        if isinstance(item, dict)
    }
    if comparison_mode == "correction":
        historical = by_label.get("historical_impact")
        repeatability = by_label.get("candidate_repeatability")
        passed = (
            not failures
            and len(comparisons) == 2
            and historical is not None
            and historical.get("status") in {"changed", "unchanged"}
            and repeatability is not None
            and repeatability.get("status") == "pass"
        )
    else:
        baseline = by_label.get("baseline_vs_candidate_1")
        repeatability = by_label.get("candidate_repeatability")
        passed = (
            not failures
            and len(comparisons) == 2
            and baseline is not None
            and baseline.get("status") == "pass"
            and repeatability is not None
            and repeatability.get("status") == "pass"
        )
    report = {
        "schema_version": 2,
        "comparison_mode": comparison_mode,
        "status": "pass" if passed else "fail",
        "input_sha256_after_runs": observed_input_hash,
        "comparisons": comparisons,
        "failures": failures,
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=report_path.parent,
            prefix=f".{report_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
            json.dump(report, handle, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, report_path)
    except OSError:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    if announce:
        print(f"Stata run comparison: {report['status'].upper()}")
    return report


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--baseline-run", required=True)
    result.add_argument("--candidate-run-1", required=True)
    result.add_argument("--candidate-run-2", required=True)
    result.add_argument("--input-file", required=True)
    result.add_argument("--analysis-root", default=str(ROOT))
    result.add_argument("--expected-input-sha256")
    result.add_argument(
        "--comparison-mode",
        choices=("equivalence", "correction"),
        default="equivalence",
    )
    result.add_argument("--report", required=True)
    return result


def main(argv: Iterable[str] | None = None) -> int:
    args = parser().parse_args(argv)
    report_path = Path(args.report)
    write_report(
        report_path,
        observed_input_hash=None,
        comparisons=[],
        failures=[{"category": "comparison_incomplete", "role": "run_set"}],
        comparison_mode=args.comparison_mode,
        announce=False,
    )

    baseline = Path(args.baseline_run).resolve()
    candidate_one = Path(args.candidate_run_1).resolve()
    candidate_two = Path(args.candidate_run_2).resolve()
    runs = {
        "baseline": baseline,
        "candidate_1": candidate_one,
        "candidate_2": candidate_two,
    }
    input_file = Path(args.input_file)
    analysis_root = Path(args.analysis_root).resolve()
    if input_file.name != input_manifest.INPUT_FILENAME:
        report = write_report(
            report_path,
            observed_input_hash=None,
            comparisons=[],
            failures=[
                {
                    "category": "input_manifest_malformed",
                    "role": "current_input",
                }
            ],
            comparison_mode=args.comparison_mode,
        )
        return 0 if report["status"] == "pass" else 1

    try:
        current_input = input_manifest.validate_approved_input(
            input_file.parent,
            analysis_root,
        )
    except input_manifest.InputManifestFailure as error:
        report = write_report(
            report_path,
            observed_input_hash=None,
            comparisons=[],
            failures=[
                {
                    "category": error.category,
                    "role": "current_input",
                }
            ],
            comparison_mode=args.comparison_mode,
        )
        return 0 if report["status"] == "pass" else 1

    try:
        loaded_runs = {role: load_run(path) for role, path in runs.items()}
        manifests = {
            role: loaded_runs[role][0]
            for role in runs
        }
        artifact_roots = {
            role: loaded_runs[role][1]
            for role in runs
        }
        failures: list[dict[str, Any]] = validate_run_set(
            runs,
            manifests,
            artifact_roots,
        )
        failures.extend(
            validate_live_run_evidence(
                runs,
                manifests,
                artifact_roots,
            )
        )
        failures.extend(
            validate_run_input_manifest_references(
                manifests,
                current_input["approval"],
            )
        )
    except ArtifactRootFailure:
        report = write_report(
            report_path,
            observed_input_hash=current_input["input"]["sha256"],
            comparisons=[],
            failures=[
                {
                    "category": "artifact_root",
                    "role": "run_set",
                }
            ],
            comparison_mode=args.comparison_mode,
        )
        return 0 if report["status"] == "pass" else 1
    except (AttributeError, KeyError, OSError, TypeError, ValueError):
        report = write_report(
            report_path,
            observed_input_hash=current_input["input"]["sha256"],
            comparisons=[],
            failures=[
                {
                    "category": "run_manifest_unavailable",
                    "role": "run_set",
                }
            ],
            comparison_mode=args.comparison_mode,
        )
        return 0 if report["status"] == "pass" else 1

    observed_input_hash = current_input["input"]["sha256"]
    failures.extend(
        validate_post_run_input(
            observed_input_hash,
            manifests,
            args.expected_input_sha256,
        )
    )

    comparisons = []
    if not failures:
        try:
            baseline_comparison = compare_pair(
                baseline,
                candidate_one,
                (
                    "historical_impact"
                    if args.comparison_mode == "correction"
                    else "baseline_vs_candidate_1"
                ),
            )
            if args.comparison_mode == "correction":
                role_map = {
                    "left": "baseline",
                    "right": "candidate_1",
                }
                for failure in baseline_comparison["failures"]:
                    if (
                        failure.get("category")
                        in CORRECTION_MODE_INTEGRITY_CATEGORIES
                    ):
                        failures.append(
                            {
                                "category": failure["category"],
                                "role": role_map.get(
                                    failure.get("role"),
                                    "run_set",
                                ),
                            }
                        )
                baseline_comparison["status"] = (
                    "unchanged"
                    if not baseline_comparison["failures"]
                    else "changed"
                )
            comparisons = [
                baseline_comparison,
                compare_pair(
                    candidate_one,
                    candidate_two,
                    "candidate_repeatability",
                ),
            ]
        except (
            AttributeError,
            KeyError,
            OSError,
            TypeError,
            ValueError,
            zipfile.BadZipFile,
        ):
            comparisons = []
            failures.append({"category": "comparison_error", "role": "run_set"})

    try:
        final_input = input_manifest.validate_approved_input(
            input_file.parent,
            analysis_root,
        )
    except input_manifest.InputManifestFailure as error:
        failures.append({"category": error.category, "role": "current_input"})
    else:
        if (
            final_input["input"] != current_input["input"]
            or final_input["approval"] != current_input["approval"]
        ):
            failures.append(
                {
                    "category": "input_manifest_comparison_drift",
                    "role": "current_input",
                }
            )

    report = write_report(
        report_path,
        observed_input_hash=observed_input_hash,
        comparisons=comparisons,
        failures=failures,
        comparison_mode=args.comparison_mode,
    )
    return 0 if report["status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
